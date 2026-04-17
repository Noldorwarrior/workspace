#!/usr/bin/env python3
"""
diff_versions.py — Diff между версиями файлов (docx, xlsx, txt).
Показывает добавленные, удалённые и изменённые элементы.
"""

import argparse, json, sys, difflib
from pathlib import Path

def extract_text(filepath):
    ext = Path(filepath).suffix.lower()
    if ext == ".docx":
        from docx import Document
        doc = Document(filepath)
        return [p.text for p in doc.paragraphs if p.text.strip()]
    elif ext in (".xlsx", ".xls"):
        import openpyxl
        wb = openpyxl.load_workbook(filepath, data_only=True)
        lines = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                lines.append("\t".join(str(c) if c else "" for c in row))
        wb.close()
        return lines
    elif ext in (".txt", ".md", ".csv"):
        return Path(filepath).read_text(encoding="utf-8").splitlines()
    return []

def verify(files):
    if len(files) < 2:
        return {"script": "diff_versions.py", "status": "skip",
                "details": "Нужны 2 файла для сравнения", "items_checked": 0,
                "items_passed": 0, "items_warned": 0, "items_failed": 0, "findings": []}

    old_lines = extract_text(files[0])
    new_lines = extract_text(files[1])
    
    differ = difflib.unified_diff(old_lines, new_lines, fromfile=files[0], tofile=files[1], lineterm="")
    diff_lines = list(differ)
    
    added = len([l for l in diff_lines if l.startswith("+") and not l.startswith("+++")])
    removed = len([l for l in diff_lines if l.startswith("-") and not l.startswith("---")])
    
    findings = []
    if added > 0 or removed > 0:
        findings.append({
            "severity": "info",
            "location": f"{files[0]} → {files[1]}",
            "expected": "—",
            "actual": f"+{added} / -{removed}",
            "description": f"Добавлено строк: {added}, удалено: {removed}",
        })

    # Первые 10 изменений для контекста
    change_count = 0
    for line in diff_lines:
        if line.startswith("+") and not line.startswith("+++"):
            if change_count < 10:
                findings.append({
                    "severity": "info",
                    "location": "Добавлено",
                    "expected": "—",
                    "actual": line[1:][:100],
                    "description": f"+ {line[1:][:100]}",
                })
            change_count += 1
        elif line.startswith("-") and not line.startswith("---"):
            if change_count < 10:
                findings.append({
                    "severity": "info",
                    "location": "Удалено",
                    "expected": line[1:][:100],
                    "actual": "—",
                    "description": f"- {line[1:][:100]}",
                })
            change_count += 1

    items_checked = max(len(old_lines), len(new_lines))
    return {
        "script": "diff_versions.py",
        "status": "info",
        "details": f"Строк: {len(old_lines)} → {len(new_lines)}, изменений: +{added}/-{removed}",
        "items_checked": items_checked,
        "items_passed": items_checked,
        "items_warned": 0,
        "items_failed": 0,
        "findings": findings,
    }

if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("files", nargs="+")
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()
    result = verify(args.files)
    print(json.dumps(result, ensure_ascii=False, indent=2) if args.json else f"📝 Diff: {result['details']}")
